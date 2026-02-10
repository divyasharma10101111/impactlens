from __future__ import annotations
import io
from datetime import datetime
from typing import List

import matplotlib.pyplot as plt

from analytics.kpis import SOVRow
from analytics.alerts import Alert
from analytics.benchmarks import BenchmarkRow
from analytics.recommendations import Recommendation

def _bar_png(values, labels, title) -> bytes:
    fig = plt.figure(figsize=(10, 4))
    ax = fig.add_subplot(111)
    ax.bar(range(len(values)), values)
    ax.set_xticks(range(len(labels)))
    ax.set_xticklabels(labels, rotation=25, ha="right")
    ax.set_title(title)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=180)
    plt.close(fig)
    buf.seek(0)
    return buf.read()

def build_client_deck_pptx(
    client_name: str,
    week_start: datetime,
    week_end: datetime,
    sov: List[SOVRow],
    top_posts,
    alerts: List[Alert],
    benchmarks: List[BenchmarkRow],
    recs: List[Recommendation],
    brand_name: str = "ImpactLens",
) -> bytes:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except Exception as e:
        raise RuntimeError("python-pptx is required. Install with: pip install python-pptx") from e

    prs = Presentation()

    # Title
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = f"Weekly Sponsor Report — {brand_name}"
    slide.placeholders[1].text = f"{client_name}\n{week_start.date()} → {week_end.date()}"

    # Summary
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Executive Summary"
    total_views = sum(r.total_views for r in sov) if sov else 0
    total_eng = sum(r.total_engagements for r in sov) if sov else 0
    total_posts = sum(r.sponsored_posts for r in sov) if sov else 0

    tf = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.8), Inches(4.8)).text_frame
    tf.word_wrap = True
    tf.text = f"Sponsored Posts: {total_posts:,}\nSponsored Views: {total_views:,}\nEngagements: {total_eng:,}"
    tf.add_paragraph().text = f"Under-delivery alerts: {len(alerts)}"
    tf.add_paragraph().text = f"Recommendations: {min(len(recs), 6)} key actions for next week"

    # SOV chart
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Sponsor Share of Voice (Views)"
    if sov:
        png = _bar_png([r.total_views for r in sov[:8]], [r.sponsor for r in sov[:8]], "Top sponsors by Views")
        slide.shapes.add_picture(io.BytesIO(png), Inches(0.8), Inches(1.5), width=Inches(8.8))
    else:
        slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8), Inches(1)).text_frame.text = "No data."

    # Alerts
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Under-delivery Alerts"
    tf = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.8), Inches(4.8)).text_frame
    tf.word_wrap = True
    if not alerts:
        tf.text = "No alerts."
    else:
        tf.text = alerts[0].message
        for a in alerts[1:8]:
            tf.add_paragraph().text = a.message

    # Benchmarks
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Benchmarks (Cohort Percentiles)"
    tf = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.8), Inches(4.8)).text_frame
    tf.word_wrap = True
    if not benchmarks:
        tf.text = "Not enough cohort data yet."
    else:
        tf.text = "Key cohort positions:"
        for b in benchmarks[:8]:
            tf.add_paragraph().text = f"{b.cohort_key} • {b.metric}: {b.percentile*100:.0f}% (you {b.client_value:,} vs median {b.cohort_median:,})"

    # Recommendations
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Next Week Recommendations"
    tf = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.8), Inches(4.8)).text_frame
    tf.word_wrap = True
    if not recs:
        tf.text = "Not enough data to generate recommendations yet."
    else:
        tf.text = "Top actions:"
        for r in recs[:6]:
            tf.add_paragraph().text = f"• {r.detail}"

    # Top posts
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Top Sponsored Creatives"
    tf = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.8), Inches(4.8)).text_frame
    tf.word_wrap = True
    if not top_posts:
        tf.text = "No top posts."
    else:
        first = top_posts[0]
        tf.text = f"1) {first.platform.upper()} • {first.media_type} • {int(first.views or 0):,} views"
        tf.add_paragraph().text = f"{first.url}"
        for i, p0 in enumerate(top_posts[1:8], start=2):
            tf.add_paragraph().text = f"{i}) {p0.platform.upper()} • {p0.media_type} • {int(p0.views or 0):,} views • {p0.url}"

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.read()
